"""Inspeção das execuções anteriores em `output/{run_id}/`."""

import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from social_media.src.config import build_config

logger = logging.getLogger(__name__)

METRICS_FILENAME = "metrics.xlsx"
CHARTS_DIRNAME = "charts"


@dataclass(frozen=True)
class RunInfo:
    """Metadados de uma execução encontrada em disco.

    Attributes:
        run_id: Nome do diretório da execução.
        path: Caminho do diretório da execução.
        created_at: Data/hora de criação do diretório.
        complete: True se os artefatos mínimos existem (metrics.xlsx e gráficos).
        posts_count: Número de posts processados (None se indisponível).
        comments_count: Número de comentários processados (None se indisponível).
        xlsx_path: Caminho do metrics.xlsx (None se ausente).
        csv_files: Arquivos CSV presentes.
        chart_files: Gráficos PNG presentes.
    """

    run_id: str
    path: Path
    created_at: datetime
    complete: bool
    posts_count: int | None
    comments_count: int | None
    xlsx_path: Path | None
    csv_files: list[Path] = field(default_factory=list)
    chart_files: list[Path] = field(default_factory=list)


def resolve_output_dir() -> Path:
    """Resolve o diretório de saída sem exigir a chave do Apify.

    Usa a configuração padrão quando disponível; caso contrário, cai no
    diretório `output/` relativo ao diretório de trabalho.
    """
    try:
        return build_config().output_dir
    except ValueError:
        return Path("output").resolve()


def _count_sheet_rows(xlsx_path: Path, sheet_name: str) -> int | None:
    """Conta as linhas de dados de uma aba do XLSX (desconta o cabeçalho)."""
    try:
        workbook = load_workbook(xlsx_path, read_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                return None
            worksheet = workbook[sheet_name]
            if worksheet.max_row is None:
                return None
            return max(worksheet.max_row - 1, 0)
        finally:
            workbook.close()
    except Exception:
        logger.warning(
            "Falha ao ler a aba '%s' do XLSX.",
            sheet_name,
            extra={"xlsx_path": str(xlsx_path)},
        )
        return None


def get_run_info(run_dir: Path) -> RunInfo:
    """Coleta os metadados de um diretório de execução."""
    run_dir = Path(run_dir)
    xlsx_path = run_dir / METRICS_FILENAME
    charts_dir = run_dir / CHARTS_DIRNAME

    has_xlsx = xlsx_path.is_file()
    chart_files = sorted(charts_dir.glob("*.png")) if charts_dir.is_dir() else []
    csv_files = sorted(run_dir.glob("*.csv"))

    posts_count: int | None = None
    comments_count: int | None = None
    if has_xlsx:
        posts_count = _count_sheet_rows(xlsx_path, "Posts")
        comments_count = _count_sheet_rows(xlsx_path, "Comentários")

    complete = has_xlsx and len(chart_files) > 0

    return RunInfo(
        run_id=run_dir.name,
        path=run_dir,
        created_at=datetime.fromtimestamp(run_dir.stat().st_ctime),
        complete=complete,
        posts_count=posts_count,
        comments_count=comments_count,
        xlsx_path=xlsx_path if has_xlsx else None,
        csv_files=csv_files,
        chart_files=chart_files,
    )


def list_runs(output_dir: Path) -> list[RunInfo]:
    """Lista todas as execuções em `output_dir`, da mais recente à mais antiga."""
    output_dir = Path(output_dir)
    if not output_dir.is_dir():
        return []

    runs = [
        get_run_info(child)
        for child in output_dir.iterdir()
        if child.is_dir()
    ]
    return sorted(runs, key=lambda run: run.created_at, reverse=True)
